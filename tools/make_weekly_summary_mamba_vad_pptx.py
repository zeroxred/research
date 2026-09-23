from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentations" / "06_weekly_summary_mamba_vad_en.pptx"

WIDE_LAYOUT = (13.333, 7.5)

COLORS = {
    "bg": RGBColor(248, 250, 252),
    "ink": RGBColor(17, 24, 39),
    "muted": RGBColor(75, 85, 99),
    "line": RGBColor(217, 119, 6),
    "blue": RGBColor(219, 234, 254),
    "blue_text": RGBColor(29, 78, 216),
    "green": RGBColor(209, 250, 229),
    "green_text": RGBColor(4, 120, 87),
    "amber": RGBColor(254, 243, 199),
    "amber_text": RGBColor(180, 83, 9),
    "sky": RGBColor(224, 242, 254),
    "sky_text": RGBColor(7, 89, 133),
    "rose": RGBColor(255, 228, 230),
    "rose_text": RGBColor(190, 18, 60),
    "gray": RGBColor(229, 231, 235),
}


def set_background(slide, color=COLORS["bg"]):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, text: str):
    box = slide.shapes.add_textbox(Inches(0.65), Inches(0.42), Inches(12.0), Inches(0.62))
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = text
    p.font.name = "Aptos Display"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = COLORS["ink"]
    return box


def add_subtitle(slide, text: str, x=0.72, y=2.65, w=10.8, h=0.5):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(20)
    p.font.color.rgb = COLORS["muted"]
    return box


def add_bullets(slide, bullets: list[str], x, y, w, h, size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = COLORS["muted"]
        p.space_after = Pt(8)
    return box


def add_text(slide, text: str, x, y, w, h, size=20, bold=False, color=None, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.text = text
    p.font.name = "Aptos Display" if bold else "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color or COLORS["ink"]
    if align is not None:
        p.alignment = align
    return box


def add_card(slide, x, y, w, h, fill_color):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = fill_color
    shape.adjustments[0] = 0.08
    return shape


def add_rule(slide, x, y, w, color=COLORS["line"]):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(0.05),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color
    return shape


def blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    return slide


def title_slide(prs):
    slide = blank_slide(prs)
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.9))
    band.fill.solid()
    band.fill.fore_color.rgb = COLORS["ink"]
    band.line.color.rgb = COLORS["ink"]
    add_text(slide, "Weekly Progress Summary", 0.72, 1.75, 11.5, 0.8, 38, True)
    add_subtitle(slide, "Mamba-based video anomaly detection", y=2.75)
    add_rule(slide, 0.72, 3.58, 2.2)
    add_bullets(
        slide,
        [
            "Started the Mamba-for-VAD block after the general Video Mamba foundation.",
            "Reviewed unsupervised VAD papers that use Mamba as an efficient temporal backbone.",
            "Connected the Mamba line back to WSVAD through multilingual prompt-guided learning.",
        ],
        0.78,
        4.05,
        11.1,
        1.5,
        21,
    )


def build(prs):
    title_slide(prs)

    slide = blank_slide(prs)
    add_title(slide, "Narrative of the Work")
    add_bullets(
        slide,
        [
            "The previous block established why Mamba is attractive for long video sequences.",
            "This block asks a more specific question: how is Mamba used in video anomaly detection?",
            "The first papers are mostly unsupervised: they learn normality and detect deviations.",
            "The final paper reconnects the discussion to weak supervision, CLIP, prompts, and temporal reasoning.",
        ],
        0.85,
        1.32,
        11.4,
        4.6,
        24,
    )

    slide = blank_slide(prs)
    add_title(slide, "Why This Step Matters")
    add_bullets(
        slide,
        [
            "VAD requires temporal memory because anomalies may unfold slowly or appear briefly.",
            "Transformers provide global reasoning but can be expensive for long surveillance videos.",
            "Mamba offers linear-time sequence modeling and a practical path for long-range temporal context.",
            "The key is deciding whether Mamba should replace, complement, or simply strengthen existing VAD pipelines.",
        ],
        0.75,
        1.25,
        5.9,
        4.8,
        21,
    )
    add_card(slide, 7.05, 1.25, 5.3, 3.9, COLORS["gray"])
    add_text(slide, "Key Idea", 7.35, 1.68, 4.5, 0.45, 27, True)
    add_text(
        slide,
        "The work shifts from understanding Mamba as a video backbone to evaluating Mamba as a VAD design choice.",
        7.35,
        2.42,
        4.55,
        1.6,
        23,
        False,
        COLORS["muted"],
    )

    slide = blank_slide(prs)
    add_title(slide, "Two VAD Directions")
    add_card(slide, 0.8, 1.45, 5.75, 4.1, COLORS["blue"])
    add_card(slide, 6.8, 1.45, 5.75, 4.1, COLORS["green"])
    add_text(slide, "Unsupervised VAD", 1.15, 1.85, 4.9, 0.45, 25, True, COLORS["blue_text"])
    add_bullets(
        slide,
        [
            "Training uses normal videos.",
            "Anomalies are detected through reconstruction or prediction errors.",
            "Typical datasets: Ped2, Avenue, ShanghaiTech.",
        ],
        1.15,
        2.55,
        4.95,
        2.2,
        19,
    )
    add_text(slide, "Weakly Supervised VAD", 7.15, 1.85, 4.9, 0.45, 25, True, COLORS["green_text"])
    add_bullets(
        slide,
        [
            "Training uses video-level labels.",
            "The model learns to localize anomalous segments.",
            "Typical datasets: UCF-Crime, XD-Violence.",
        ],
        7.15,
        2.55,
        4.95,
        2.2,
        19,
    )

    slide = blank_slide(prs)
    add_title(slide, "VADMamba")
    add_bullets(
        slide,
        [
            "VADMamba applies Mamba to reconstruction- and prediction-based VAD.",
            "It is trained only on normal videos, not weak video-level anomaly labels.",
            "The model uses frame prediction and optical-flow reconstruction as proxy tasks.",
            "At inference time, larger reconstruction or prediction errors become anomaly evidence.",
        ],
        0.75,
        1.25,
        5.55,
        4.7,
        20,
    )
    add_card(slide, 6.75, 1.45, 5.2, 3.8, COLORS["amber"])
    add_text(slide, "Contribution", 7.1, 1.85, 4.5, 0.45, 26, True, COLORS["amber_text"])
    add_text(
        slide,
        "It shows that Mamba can replace CNN or Transformer temporal modules in efficient unsupervised VAD.",
        7.1,
        2.65,
        4.45,
        1.45,
        22,
        False,
        COLORS["muted"],
    )

    slide = blank_slide(prs)
    add_title(slide, "VADMamba: What It Does Not Solve")
    add_bullets(
        slide,
        [
            "It does not compete directly with Sultani, RTFM, VADCLIP, or other WSVAD methods.",
            "It does not use CLIP-style semantic guidance.",
            "It does not learn from video-level abnormal labels.",
            "For WSVAD research, its value is mainly as evidence that Mamba is a useful temporal backbone.",
        ],
        0.85,
        1.32,
        11.4,
        4.6,
        24,
    )

    slide = blank_slide(prs)
    add_title(slide, "STNMamba")
    add_bullets(
        slide,
        [
            "STNMamba argues that vanilla Mamba is not enough for spatial-temporal normality learning.",
            "It introduces a multi-scale spatial encoder for anomalies of different sizes.",
            "It uses RGB frame differences instead of optical flow for lower-cost motion cues.",
            "It performs multi-level spatial-temporal fusion and stores normal prototypes in memory banks.",
        ],
        0.75,
        1.25,
        5.65,
        4.7,
        20,
    )
    add_card(slide, 6.9, 1.35, 4.9, 0.82, COLORS["blue"])
    add_card(slide, 6.9, 2.45, 4.9, 0.82, COLORS["green"])
    add_card(slide, 6.9, 3.55, 4.9, 0.82, COLORS["amber"])
    add_card(slide, 6.9, 4.65, 4.9, 0.82, COLORS["sky"])
    add_text(slide, "MS-VSSB", 7.25, 1.56, 4.2, 0.35, 20, True, COLORS["blue_text"], PP_ALIGN.CENTER)
    add_text(slide, "CA-VSSB", 7.25, 2.66, 4.2, 0.35, 20, True, COLORS["green_text"], PP_ALIGN.CENTER)
    add_text(slide, "STIM + STFB", 7.25, 3.76, 4.2, 0.35, 20, True, COLORS["amber_text"], PP_ALIGN.CENTER)
    add_text(slide, "Memory Banks", 7.25, 4.86, 4.2, 0.35, 20, True, COLORS["sky_text"], PP_ALIGN.CENTER)

    slide = blank_slide(prs)
    add_title(slide, "M2S2L")
    add_bullets(
        slide,
        [
            "M2S2L extends Mamba-based VAD through explicit multi-scale learning.",
            "Spatial scales capture small local changes, medium patterns, and larger scene structure.",
            "Temporal scales capture short motion, medium dynamics, and long behavior.",
            "Feature decomposition separates common, appearance-specific, and motion-specific representations.",
        ],
        0.75,
        1.25,
        5.75,
        4.7,
        20,
    )
    add_card(slide, 7.0, 1.45, 4.9, 1.0, COLORS["blue"])
    add_card(slide, 7.0, 2.95, 4.9, 1.0, COLORS["green"])
    add_card(slide, 7.0, 4.45, 4.9, 1.0, COLORS["rose"])
    add_text(slide, "Multi-scale spatial", 7.28, 1.76, 4.3, 0.35, 21, True, COLORS["blue_text"], PP_ALIGN.CENTER)
    add_text(slide, "Multi-scale temporal", 7.28, 3.26, 4.3, 0.35, 21, True, COLORS["green_text"], PP_ALIGN.CENTER)
    add_text(slide, "Feature decomposition", 7.28, 4.76, 4.3, 0.35, 21, True, COLORS["rose_text"], PP_ALIGN.CENTER)

    slide = blank_slide(prs)
    add_title(slide, "VADMamba++")
    add_bullets(
        slide,
        [
            "VADMamba++ changes the proxy task rather than only changing the architecture.",
            "It replaces RGB-to-RGB reconstruction with Gray-to-RGB reasoning.",
            "The model receives grayscale frames and predicts plausible RGB frames.",
            "This removes optical flow and converts the pipeline into a single-task framework.",
        ],
        0.75,
        1.25,
        5.65,
        4.7,
        20,
    )
    add_card(slide, 6.9, 1.45, 5.1, 3.65, COLORS["gray"])
    add_text(slide, "Gray-to-RGB", 7.25, 1.9, 4.4, 0.45, 28, True)
    add_text(
        slide,
        "The anomaly signal comes from both structural geometry errors and chromatic appearance errors.",
        7.25,
        2.75,
        4.35,
        1.4,
        22,
        False,
        COLORS["muted"],
    )

    slide = blank_slide(prs)
    add_title(slide, "VADMamba++: Hybrid Modeling")
    add_card(slide, 0.85, 1.55, 3.55, 3.45, COLORS["blue"])
    add_card(slide, 4.9, 1.55, 3.55, 3.45, COLORS["green"])
    add_card(slide, 8.95, 1.55, 3.55, 3.45, COLORS["amber"])
    add_text(slide, "Transformer", 1.15, 1.95, 2.95, 0.45, 24, True, COLORS["blue_text"])
    add_text(slide, "Global context modeling.", 1.15, 2.72, 2.95, 0.9, 21, False, COLORS["muted"])
    add_text(slide, "Mamba", 5.2, 1.95, 2.95, 0.45, 24, True, COLORS["green_text"])
    add_text(slide, "Long-range temporal dependencies.", 5.2, 2.72, 2.95, 0.9, 21, False, COLORS["muted"])
    add_text(slide, "CNN", 9.25, 1.95, 2.95, 0.45, 24, True, COLORS["amber_text"])
    add_text(slide, "Local spatial refinement.", 9.25, 2.72, 2.95, 0.9, 21, False, COLORS["muted"])
    add_bullets(
        slide,
        [
            "The best encoder order is coarse-to-fine: Transformer, then Mamba, then CNN.",
        ],
        1.0,
        5.55,
        11.1,
        0.7,
        20,
    )

    slide = blank_slide(prs)
    add_title(slide, "Results Pattern")
    add_bullets(
        slide,
        [
            "VADMamba proves that selective SSMs can be effective for unsupervised VAD.",
            "STNMamba improves spatial-temporal learning while remaining lightweight.",
            "M2S2L improves accuracy by specializing scale and modality representations.",
            "VADMamba++ improves the accuracy-speed trade-off with Gray-to-RGB and no optical flow.",
        ],
        0.85,
        1.25,
        11.4,
        3.2,
        22,
    )
    add_card(slide, 1.0, 4.85, 2.5, 0.9, COLORS["blue"])
    add_card(slide, 3.9, 4.85, 2.5, 0.9, COLORS["green"])
    add_card(slide, 6.8, 4.85, 2.5, 0.9, COLORS["amber"])
    add_card(slide, 9.7, 4.85, 2.5, 0.9, COLORS["rose"])
    add_text(slide, "Backbone", 1.2, 5.12, 2.1, 0.3, 18, True, COLORS["blue_text"], PP_ALIGN.CENTER)
    add_text(slide, "Architecture", 4.1, 5.12, 2.1, 0.3, 18, True, COLORS["green_text"], PP_ALIGN.CENTER)
    add_text(slide, "Scale", 7.0, 5.12, 2.1, 0.3, 18, True, COLORS["amber_text"], PP_ALIGN.CENTER)
    add_text(slide, "Proxy task", 9.9, 5.12, 2.1, 0.3, 18, True, COLORS["rose_text"], PP_ALIGN.CENTER)

    slide = blank_slide(prs)
    add_title(slide, "Multilingual Prompt-Guided WSVAD")
    add_bullets(
        slide,
        [
            "The final paper returns to the weakly supervised setting.",
            "It extends CLIP-style WSVAD with multilingual prompt pools.",
            "Adaptive Top-K selection chooses the most useful prompts for each video snippet.",
            "Mamba and Transformer are combined for multi-granularity temporal modeling.",
        ],
        0.75,
        1.25,
        5.65,
        4.7,
        20,
    )
    add_card(slide, 6.9, 1.45, 5.0, 3.8, COLORS["sky"])
    add_text(slide, "Bridge to WSVAD", 7.25, 1.9, 4.3, 0.45, 27, True, COLORS["sky_text"])
    add_text(
        slide,
        "Here Mamba is not the whole VAD paradigm. It is a temporal reasoning module inside a CLIP and prompt-guided pipeline.",
        7.25,
        2.7,
        4.25,
        1.55,
        21,
        False,
        COLORS["muted"],
    )

    slide = blank_slide(prs)
    add_title(slide, "Evolution of Ideas")
    add_card(slide, 0.45, 2.05, 2.0, 1.25, COLORS["blue"])
    add_card(slide, 2.82, 2.05, 2.0, 1.25, COLORS["green"])
    add_card(slide, 5.19, 2.05, 2.0, 1.25, COLORS["amber"])
    add_card(slide, 7.56, 2.05, 2.0, 1.25, COLORS["sky"])
    add_card(slide, 9.93, 2.05, 2.95, 1.25, COLORS["rose"])
    add_text(slide, "VADMamba", 0.65, 2.43, 1.6, 0.4, 18, True, COLORS["blue_text"], PP_ALIGN.CENTER)
    add_text(slide, "STNMamba", 3.02, 2.43, 1.6, 0.4, 18, True, COLORS["green_text"], PP_ALIGN.CENTER)
    add_text(slide, "M2S2L", 5.39, 2.43, 1.6, 0.4, 18, True, COLORS["amber_text"], PP_ALIGN.CENTER)
    add_text(slide, "VADM++", 7.76, 2.43, 1.6, 0.4, 18, True, COLORS["sky_text"], PP_ALIGN.CENTER)
    add_text(slide, "Multilingual WSVAD", 10.08, 2.43, 2.65, 0.4, 17, True, COLORS["rose_text"], PP_ALIGN.CENTER)
    add_bullets(
        slide,
        [
            "From Mamba as an efficient temporal backbone.",
            "To spatial-temporal normality learning with memory.",
            "To multi-scale and modality-specialized reconstruction.",
            "To simpler proxy tasks and better deployment efficiency.",
            "To CLIP, multilingual prompts, and weakly supervised localization.",
        ],
        1.05,
        4.05,
        11.3,
        1.85,
        21,
    )

    slide = blank_slide(prs)
    add_title(slide, "What Is Already in Place")
    add_bullets(
        slide,
        [
            "A clear distinction between unsupervised VAD and WSVAD evaluation protocols.",
            "A map of how Mamba enters VAD: backbone, spatial-temporal fusion, memory, and prompt-guided temporal modeling.",
            "A comparison of architectural improvements versus proxy-task improvements.",
            "A bridge from reconstruction-based Mamba papers back to the main weakly supervised research direction.",
        ],
        0.85,
        1.3,
        11.4,
        4.8,
        22,
    )

    slide = blank_slide(prs)
    add_title(slide, "Suggested Next Step")
    add_bullets(
        slide,
        [
            "Build a comparison matrix separating unsupervised VAD from WSVAD methods.",
            "For each method, record supervision type, backbone, temporal module, anomaly score, datasets, and metrics.",
            "Decide whether the PhD direction should use Mamba as a full backbone or as a temporal module inside a CLIP/WSVAD pipeline.",
            "Use that decision to define the first implementation target and baseline set.",
        ],
        0.85,
        1.3,
        11.4,
        4.8,
        22,
    )


def main():
    prs = Presentation()
    prs.slide_width = Inches(WIDE_LAYOUT[0])
    prs.slide_height = Inches(WIDE_LAYOUT[1])
    build(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)
    print(f"{len(prs.slides)} slides")


if __name__ == "__main__":
    main()
